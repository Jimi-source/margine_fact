#!/usr/bin/env python3
"""Собирает xlsx-вид из файла, который отдан аналитику, — чтобы посмотреть глазами.

    python3 tools/handover_to_xlsx.py [путь_к_json] [путь_к_xlsx]

По умолчанию читает ~/Desktop/Ozon/margine_fact_data.json и пишет
данные_для_аналитика.xlsx в корень проекта. Ничего не добавляет и не считает —
это ровно то же содержимое, просто в виде таблицы.
"""
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HOME = os.path.expanduser("~")
SRC = sys.argv[1] if len(sys.argv) > 1 else f"{HOME}/Desktop/Ozon/margine_fact_data.json"
OUT = sys.argv[2] if len(sys.argv) > 2 else f"{HOME}/Desktop/margine_fact/данные_для_аналитика.xlsx"

d = json.load(open(SRC, encoding="utf-8"))

F = "Arial"
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E79")
BOLD = Font(name=F, size=10, bold=True)
BASE = Font(name=F, size=10)
MONEY = "#,##0;(#,##0);-"
PCT = "0.0%;(0.0%);-"
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

wb = Workbook()


def header(ws, cols, widths=None):
    for i, t in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=t)
        c.font = HEAD
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(t, 13)
    ws.row_dimensions[1].height = 32


TEXT_COLS = {"период", "артикул", "шт", "хвост_недель"}

ws = wb.active
ws.title = "Неделя × SKU"
COLS = ["период", "хвост_недель", "доля_доехавших_%", "артикул", "шт", "продажи", "возвраты",
        "вознаграждение_ozon", "услуги_доставки", "услуги_партнёров", "реклама_из_начислений",
        "другие_услуги_и_штрафы", "начисления_итого", "реклама_трафарет", "реклама_пвп",
        "прочие_услуги_разнесённые", "себестоимость", "сумма_отмен", "выручка", "маржа_доля"]
header(ws, COLS, {"период": 21, "артикул": 21, "хвост_недель": 9, "доля_доехавших_%": 11,
                  "шт": 6, "маржа_доля": 10})
r = 2
for row in d["byWeek"]:
    for i, c in enumerate(COLS, start=1):
        cell = ws.cell(row=r, column=i, value=row.get(c))
        cell.font = BASE
        cell.border = THIN
        if c == "маржа_доля":
            cell.number_format = PCT
        elif c == "доля_доехавших_%":
            cell.number_format = "0.00"
        elif c not in TEXT_COLS:
            cell.number_format = MONEY
    r += 1
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{r - 1}"

ws2 = wb.create_sheet("Итоги по неделям")
C2 = ["период", "хвост_недель", "доля_доехавших_%", "шт", "продажи", "начисления", "реклама", "себес", "прочие"]
header(ws2, C2, {"период": 21, "хвост_недель": 9, "доля_доехавших_%": 11, "шт": 7})
r = 2
for w in d["perWeek"]:
    for i, c in enumerate(C2, start=1):
        cell = ws2.cell(row=r, column=i, value=w.get(c))
        cell.font = BASE
        cell.border = THIN
        if c == "доля_доехавших_%":
            cell.number_format = "0.00"
        elif c not in TEXT_COLS:
            cell.number_format = MONEY
    r += 1

ws3 = wb.create_sheet("Поля и источник")
ws3.column_dimensions["A"].width = 30
ws3.column_dimensions["B"].width = 105
ws3["A1"] = "Что получил аналитик"
ws3["A1"].font = Font(name=F, size=13, bold=True)
ws3["A2"] = f"Файл: {SRC}"
ws3["A2"].font = Font(name=F, size=9, italic=True, color="808080")

meta = d["meta"]
r = 4
ws3.cell(row=r, column=1, value="ВЫГРУЗКА").font = BOLD
r += 1
for k in ["выгружено", "пользователь", "период", "недель_взято", "последний_рассчитанный_период"]:
    if k in meta:
        ws3.cell(row=r, column=1, value=k).font = BASE
        c = ws3.cell(row=r, column=2, value=str(meta[k]))
        c.font = BASE
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
if "как_читать_хвост" in meta:
    ws3.cell(row=r, column=1, value="как_читать_хвост").font = BASE
    c = ws3.cell(row=r, column=2, value=meta["как_читать_хвост"])
    c.font = BASE
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.row_dimensions[r].height = 60
    r += 2

ws3.cell(row=r, column=1, value="ПОЛЯ").font = BOLD
r += 1
for k, v in d["_поля"].items():
    ws3.cell(row=r, column=1, value=k).font = BASE
    ws3.cell(row=r, column=1).alignment = Alignment(vertical="top")
    c = ws3.cell(row=r, column=2, value=v)
    c.font = BASE
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.row_dimensions[r].height = 28 if len(v) > 100 else 14
    r += 1

r += 1
ws3.cell(row=r, column=1, value="ХВОСТ ДОСТАВКИ").font = BOLD
r += 1
ws3.cell(row=r, column=1, value="лаг, недель").font = BOLD
ws3.cell(row=r, column=2, value="накоплено % начислений").font = BOLD
r += 1
for x in meta.get("хвост_доставки", []):
    ws3.cell(row=r, column=1, value=x["недель"]).font = BASE
    ws3.cell(row=r, column=2, value=x["накоплено"]).font = BASE
    r += 1

wb.save(OUT)
print("saved:", OUT)
print(f"строк неделя×SKU: {len(d['byWeek'])} | недель: {len(d['perWeek'])} | выгружено: {meta.get('выгружено')}")
